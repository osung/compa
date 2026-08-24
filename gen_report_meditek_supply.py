# -*- coding: utf-8 -*-
"""2026 MEDITEK 기업 기술 × 공급기관 수행 국가 R&D 과제 매칭 보고서(docx).

gen_report_meditek_top10.py 의 형제 판. 서식·과제 상세 블록은 gen_report 를, 문구 정리
헬퍼(neutralize/rename_rnd/sort_key)는 gen_report_meditek_top10 을 그대로 재사용하고
아래가 다르다.

  - 입력이 MEDITEK_260820_보고서.json (match_meditek_supply.py 산출, top5 스키마)
  - 추천 표에 '공급기관'(기술이전 창구) 열을 넣는다 — 매칭 대상이 공급기관 수행 과제이므로
    수요기업이 실제로 접촉할 기관을 표에서 바로 읽을 수 있어야 한다
  - '공급기관별 추천 현황' 요약 표를 대상 기업 목록 뒤에 넣는다
  - 기업 기술 블록은 합본 본문의 [라벨] 섹션을 각각 별도 행으로 편다(한 셀에 2천 자
    덩어리로 들어가면 읽히지 않는다)

※ 기업 소개면에서는 '수요기술(확보 희망)' 과 '보유기술(이미 보유)' 을 구분해서 싣는다.
  두 정보의 성격이 정반대라 뭉쳐 놓으면 읽는 사람이 보유 여부를 오해한다. 이전 판에서는
  매칭 기준을 감추려고 라벨을 없앴는데, 그 결과 기업 정보를 이해할 수 없게 되어 되돌렸다.
  단, 추천 근거 문장(판단근거·추천근거_상세)에는 여전히 '수요기술/기보유기술' 같은
  매칭 기준 용어를 쓰지 않는다(explain_meditek_top10 의 BANNED 규칙) — 근거는 기업 기술과
  과제의 접점을 설명해야 하고, 어느 쪽을 기준으로 골랐는지는 서술에 필요하지 않다.
  적합도 점수·유망성 점수는 보고서에 싣지 않는다(순위로만 제시).

입력:
  MEDITEK_260820_보고서.json      기업별 Top5 + 근거   ← match_meditek_supply.py
  $COMPA_SCRATCH/pid_fields.json  연구개발단계/연구수행주체 등 ← _proj_meta_prep.py
  $COMPA_SCRATCH/pid_patents.json 특허 실적            ← _patent_prep_nice.py

사용: COMPA_SCRATCH=<scratch> python gen_report_meditek_supply.py
"""
import glob
import json
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

import gen_report as gr
import gen_report_meditek_top10 as gm
import meditek_cite_mark as _cite
import meditek_theme as mt
import report_text_fix as tf

mt.apply_docx(gr)                      # 표·제목 색을 MEDITEK CI 로 교체(로드 직후 1회)
from gen_report import (SERIF, add_disclaimer_box, cell_vcenter, fill_cell, keep_next,
                        set_cell_width,
                        para, para_border, run_shade, section_label, set_margins,
                        set_pgnum_start, setup_cover_header, setup_running_header,
                        setup_styles, shade_cell, style_run, table_cellmar,
                        table_fixed, table_grid, rows_cantsplit, year_lines, _text_w)

_C = mt.DOCX_COLORS                    # 테마가 적용한 색을 한 곳에서 받아 쓴다
ACCENT, NAVY, INK, MUTED = _C["ACCENT"], _C["NAVY"], _C["INK"], _C["MUTED"]
HAIR, HEAD_BG, HEAD_FG = _C["HAIR"], _C["HEAD_BG"], _C["HEAD_FG"]
LABEL_BG, ZEBRA = _C["LABEL_BG"], _C["ZEBRA"]

HERE = os.path.dirname(os.path.abspath(__file__))
BEST_JSON = os.environ.get("MEDITEK_SUPPLY_JSON",
                           os.path.join(HERE, "MEDITEK_260820_보고서.json"))
OUT_BASE = "MEDITEK_공급기관_국가RnD_매칭보고서"
PUBLISH_DATE = os.environ.get("MEDITEK_PUBLISH_DATE", "2026. 8. 22.")

COVER_EYEBROW = "의료기기 / 헬스케어  OPEN INNOVATION  &  BIZ PARTNERING"
COVER_TITLE1 = "2026 MEDITEK"
COVER_TITLE2 = "국가R&D 과제 매칭 보고서"
# 표지에는 일시·장소를 싣지 않는다(요청). 행사 개요 표에만 둔다.
# 생성 주체 표기 — 발행일 아래 줄에 따로, 강조해서 싣는다
ENGINE_NOTE = "APOLLO 국가R&D 사업화 유망성 탐색 플랫폼 AI 매칭 결과"
# 표지 우측 하단 버전 표기 — 발행일 박스를 없애고 여기로 옮겼다(인쇄본 판 구분용).
# MEDITEK_REPORT_VERSION 을 주면 발행일 뒤에 판 번호가 붙는다(예: "2026. 8. 22. · v2").
_VER = os.environ.get("MEDITEK_REPORT_VERSION", "").strip()
COVER_VERSION = f"{PUBLISH_DATE} · {_VER}" if _VER else PUBLISH_DATE

# 코퍼스 조건은 '제출년도 2020 이후'로 걸지만, 제출년도는 보고서를 낸 연도라 그 해에도
# 과제가 살아 있었다는 뜻이다. 실제 매칭 결과 114개 과제의 종료연도는 2020~2030 이고
# 2020년 이전에 종료된 과제는 0건이므로(시작연도는 2013년까지 올라간다), 독자에게는
# '2020년 이후 수행 중이거나 종료된 과제'로 적는 것이 사실에 맞다.
# 특허 조건은 '등록 또는 출원 계류(공개) 특허 1건 이상'이다. 거절·취하·포기된 특허는
# 권리가 성립하지 않았고 소멸된 특허는 권리가 끝났으므로 코퍼스에서 배제한다
# (match_meditek_supply.ALIVE_STATES). "특허 1건 이상"으로만 적으면 거절된 특허만 가진
# 과제도 포함되는 것처럼 읽히므로 상태를 명시한다.
CORPUS_NOTE = ("매칭 대상 과제는 2026 MEDITEK 공급기관이 수행한 국가 R&D 과제 중 "
               "2020년 이후 수행 중이거나 종료된 과제로 한정하고, 그 가운데 등록되었거나 "
               "출원 계류 중인 특허를 1건 이상 확보한 과제만을 후보로 삼았다. 거절·취하·"
               "소멸된 특허만 있는 과제는 제외하였다. 기술이전 협의가 가능한 권리가 "
               "실제로 존재하는 과제만 추천하기 위한 조건이다.")
METHOD_NOTE = ("매칭은 기업이 제출한 기술 정보에서 핵심 기술 키워드를 추출하고, 그 키워드와 "
               "과제 임베딩의 의미 유사도, 과제가 확보한 특허와의 의미 유사도, 과제의 "
               "기술적 유망성을 함께 반영해 후보를 선별한 뒤, 각 후보가 해당 기업의 기술과 "
               "실질적으로 부합·기여하는 정도를 평가해 순위를 다시 정하는 방식으로 "
               "수행하였다. 기업의 기술과 유사한 특허를 여러 건 확보한 과제, 권리가 살아 "
               "있는 특허를 확보한 과제를 우선한다.")

TOPN_LABEL = "Top 5"
TOP_KEY = "top5"

# 기업 기술 합본의 [라벨] → 표에 쓸 행 이름. 여기 없는 라벨은 그대로 쓴다.
BODY_LABEL = {"제품(기술)요약": "기술 요약", "사업내용": "사업 내용"}
# 표의 다른 행과 겹치는 섹션은 본문에서 뺀다(산업분류=기술분야 행, 기업 키워드=핵심 키워드 행)
BODY_DROP = {"산업분류", "기업 키워드"}

_SECTION = re.compile(r"^\[([^\]\n]+)\]\n(.*?)(?=\n\n\[[^\]\n]+\]\n|\Z)", re.DOTALL | re.M)


def next_out_path():
    vers = [int(m.group(1)) for p in glob.glob(os.path.join(HERE, OUT_BASE + "_v*.docx"))
            if (m := re.search(r"_v(\d+)\.docx$", p))]
    return os.path.join(HERE, f"{OUT_BASE}_v{(max(vers) + 1) if vers else 1}.docx")


OUT = os.environ.get("MEDITEK_SUPPLY_OUT", next_out_path())


def tech_names(dm):
    """기업 기술명 — 매칭 기준 구분 없이 확보된 기술명을 그대로 나열."""
    return [x for x in (dm.get("수요기술명", ""), dm.get("기보유기술명", "")) if x.strip()]


def tech_rows(dm):
    """기업 정보표의 기술 관련 행 [(구분, 행이름, 본문, 프로즈여부), …].

    구분은 '수요기술'(확보 희망) / '보유기술'(이미 보유) 둘 중 하나다. 합본 본문이
    '[라벨]\\n본문' 섹션으로 되어 있으면 섹션별로 행을 나눈다(2천 자 한 셀 방지).
    기술명과 같은 문장인 본문은 같은 문장이 두 행에 반복되므로 싣지 않는다.
    """
    def norm(x):
        return re.sub(r"\s+", "", x or "")

    names = {norm(x) for x in tech_names(dm)}
    rows = []
    for gubun, raw in (("수요기술", dm.get("수요기술 내용", "")),
                       ("보유기술", dm.get("기보유기술 내용", ""))):
        raw = (raw or "").strip()
        if not raw:
            continue
        secs = _SECTION.findall(raw)
        if not secs:
            if norm(raw) not in names:
                rows.append((gubun, "기술 내용", raw, True))
            continue
        for label, body in secs:
            label, body = label.strip(), body.strip()
            if label in BODY_DROP or not body or norm(body) in names:
                continue
            rows.append((gubun, BODY_LABEL.get(label, label), body, True))
    return rows


def tech_name_rows(dm):
    """기술명 행 [(구분, 기술명, 구분메모)]. 수요기술과 보유기술을 따로 싣는다."""
    out = []
    if (dm.get("수요기술명") or "").strip():
        out.append(("수요기술", dm["수요기술명"], ""))
    if (dm.get("기보유기술명") or "").strip():
        gubun = " / ".join(x for x in (dm.get("기술유형", ""), dm.get("기술분야", "")) if x)
        out.append(("보유기술", dm["기보유기술명"], gubun))
    return out


# 연계유형 표기 — 표에는 짧게, 개요에 정의를 둔다.
# 유형은 기술도입·공동연구 둘뿐이다. 기업이 과제에 부품·서비스를 제공하는 방향은
# 이 보고서가 다루는 기술이전이 아니라서 유형으로 두지 않는다.
# 연계유형(기술도입/공동연구)은 **보고서에 노출하지 않는다**(요청). 매칭 내부에서
# 재랭킹 가중과 근거문 방향 규칙에만 쓰고, 산출 pkl/xlsx 에는 남긴다.
KIND_SHORT = {"기술도입": "기술도입", "공동연구": "공동연구", "미분류": "-"}
KIND_COLOR = {"기술도입": mt.APOLLO_BLUE, "공동연구": mt.APOLLO_BLUE_LT,
              "미분류": mt.APOLLO_GREY}
KIND_LEGEND = [
    ("기술도입", "공급기관 과제의 기술을 기업이 이전받아 자사 제품·서비스에 적용"),
    ("공동연구", "양측 기술을 결합해 기업과 공급기관이 함께 개발"),
]

# 구분 머리행 문구 — 두 정보의 성격을 한 번에 알 수 있게 괄호로 덧붙인다
GUBUN_LABEL = {"수요기술": "수요기술  (확보를 희망하는 기술)",
               "보유기술": "보유기술  (기업이 이미 보유한 기술·사업내용)"}
GUBUN_COLOR = {"수요기술": ACCENT, "보유기술": NAVY}   # 나비효과 두 색


def industry_row(dm):
    """기술명이 없는 기업(기업DB 근거)은 산업분류를 별도 행으로 보여 준다."""
    if tech_names(dm) or not (dm.get("기술분야") or "").strip():
        return None
    return ("산업분류", dm["기술분야"], False)


def patent_count(tp):
    """표에 싣는 특허 건수 — 상세 페이지의 특허 실적 목록과 같은 소스에서 센다."""
    return len(gr.PATENTS.get(str(tp["과제고유번호"]), []))


def supply_summary(demands):
    """공급기관별 (추천 건수, 과제 수, 특허 건수) — 요약 표용.

    특허 건수는 '과제 단위'로 센다. 같은 과제가 여러 기업에 추천되므로 추천 행마다 더하면
    본문 특허 실적 표의 합계(과제별 1회)와 어긋난다.
    """
    agg = {}
    for v in demands.values():
        for t in v[TOP_KEY]:
            sup = t.get("공급기관") or "(미상)"
            a = agg.setdefault(sup, {"rec": 0, "pids": set()})
            a["rec"] += 1
            a["pids"].add(str(t["과제고유번호"]))
    for a in agg.values():
        a["pat"] = sum(len(gr.PATENTS.get(p, [])) for p in a["pids"])
    return sorted(agg.items(), key=lambda kv: (-kv[1]["rec"], kv[0]))


def load_data():
    with open(BEST_JSON, encoding="utf-8") as f:
        demands = json.load(f)
    pf = os.path.join(gr.SCRATCH, "pid_fields.json")
    if not os.path.exists(pf):
        raise SystemExit(f"없는 파일: {pf} — 먼저 `python _proj_meta_prep.py` 실행 필요")
    with open(pf, encoding="utf-8") as f:
        pidf = json.load(f)
    pidf = pidf.get("fields", pidf)         # {pid: {...}} 로 벗겨서 넘겨야 값이 들어간다
    if not gr.PATENTS:
        raise SystemExit(f"특허 데이터 없음({gr.SCRATCH}/pid_patents.json) — "
                         f"COMPA_REPORT_JSON={BEST_JSON} 로 `python _patent_prep_nice.py` 필요")
    # 제출 원문의 표기 오류(hwp 줄바꿈 잔재로 쪼개진 단어·오탈자)를 렌더링 전에 교정한다.
    # 여기 한 곳에서 처리하면 표·상세근거·특허 실적 표가 모두 같은 교정본을 쓴다.
    log = tf.fix_report(demands, gr.PATENTS, top_key=TOP_KEY)
    print(f"[표기 교정] {tf.report_stats()}")
    for ctx, why in log[:8]:
        print(f"    {ctx}  [{why}]")
    if len(log) > 8:
        print(f"    … 그 외 {len(log) - 8}곳")
    preflight(demands, pidf)
    return demands, pidf


def preflight(demands, pidf):
    """정보표·근거·특허가 통째로 비는 사고 방지(보고서는 값이 없으면 조용히 '-' 로 찍힌다)."""
    tops = [t for v in demands.values() for t in v[TOP_KEY]]
    pids = sorted({t["과제고유번호"] for t in tops})
    leak = [k for t in tops for k in t if "유망" in k]
    leak += [k for lst in gr.PATENTS.values() for x in lst for k in x if "유망" in k]
    if leak:
        raise SystemExit(f"[preflight] 유망성 정보가 보고서 입력에 남아 있습니다: {set(leak)}")
    if not (set(pids) & set(pidf)):
        raise SystemExit("[preflight] pid_fields 에서 과제를 하나도 찾지 못했습니다 — "
                         "구조가 {pid: {...}} 인지 확인('fields' 로 감싸져 있으면 벗겨야 함).")
    print(f"[preflight] 기업 {len(demands)}건 · 추천 {len(tops)}건 · 과제 {len(pids)}개")
    for key, label in (("연구개발단계", "연구개발단계"), ("연구수행주체", "연구수행주체"),
                       ("연구책임자명", "연구책임자"), ("국가연구자번호", "국가연구자번호")):
        c = sum(1 for p in pids if str(pidf.get(p, {}).get(key, "") or "").strip())
        print(f"  {'OK ' if c == len(pids) else ('!! ' if c == 0 else '~  ')}{label:<12} "
              f"{c}/{len(pids)}")
    for label, fn in (("과제수행기간", lambda t: gr.fmt_period(t.get("과제설명문", ""))),
                      ("과학기술표준분류(중)",
                       lambda t: (gr.extract_class(t.get("과제설명문", ""))
                                  or pidf.get(t["과제고유번호"], {}).get("표준분류중"))),
                      ("과제수행기관", lambda t: t.get("수행기관", "")),
                      ("공급기관", lambda t: t.get("공급기관", "")),
                      ("매칭 근거", lambda t: t.get("판단근거", "")),
                      ("상세 매칭 근거", lambda t: t.get("추천근거_상세", ""))):
        c = sum(1 for t in tops if str(fn(t) or "").strip())
        print(f"  {'OK ' if c == len(tops) else ('!! ' if c == 0 else '~  ')}{label:<12} "
              f"{c}/{len(tops)}")
    nopat = [p for p in pids if not gr.PATENTS.get(p)]
    print(f"  {'OK ' if not nopat else '!! '}특허 실적       "
          f"{len(pids) - len(nopat)}/{len(pids)} 과제 보유"
          f" (총 {sum(len(gr.PATENTS.get(p, [])) for p in pids)}건)")
    if nopat:
        raise SystemExit(f"[preflight] 특허 성과가 없는 과제가 추천에 있습니다: {nopat[:5]} "
                         "— 매칭 조건(특허 1건 이상)과 어긋납니다")


# ---- 표지 -------------------------------------------------------------------
# PDF 판(gen_report_pdf_meditek_supply.cover)과 같은 구성을 python-docx 로 만든다.
# 캔버스가 없으므로 색 밴드·표제 패널은 음영을 준 표로 대신한다.
COVER_APOLLO_IN = 3.0          # 표지 하단 APOLLO 로고 폭(inch) ≈ 76mm, PDF 와 같은 크기
BAND_H_PT = 7                  # 색 밴드 높이


def _text_width_tw(doc):
    """구역의 실제 본문 폭(twips) — 페이지 크기를 하드코딩하지 않는다."""
    sec = doc.sections[0]
    return int((sec.page_width - sec.left_margin - sec.right_margin) / 635)


def _band(doc, cols, h=BAND_H_PT, after=0):
    """색 밴드 한 줄 — 셀마다 음영을 준 무테 표."""
    t = doc.add_table(rows=1, cols=len(cols))
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tw = _text_width_tw(doc)
    w = tw // len(cols)
    table_fixed(t, [w] * len(cols))
    for i, c in enumerate(cols):
        cell = t.rows[0].cells[i]
        shade_cell(cell, c)
        pf = cell.paragraphs[0].paragraph_format
        pf.space_before = Pt(0); pf.space_after = Pt(0)
        style_run(cell.paragraphs[0].add_run(" "), 1)
    t.rows[0].height = Pt(h)
    para(doc, "", after=after)
    return t


def build_cover(doc):                      # 건수 요약은 표지에서 빼고 개요에만 둔다
    sec = doc.sections[0]                  # 표지는 헤더가 없어 위아래 여백을 좁게 쓴다
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.45)
    para(doc, "", after=10)
    lp = doc.add_paragraph()                      # MEDITEK 공식 로고(CI 원본)
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.space_after = Pt(8)
    try:
        lp.add_run().add_picture(mt.LOGO_H, width=Inches(1.95))
    except Exception as e:
        print("로고 삽입 실패(로고 없이 진행):", e)
    para(doc, COVER_EYEBROW, 9.5, bold=True, color=mt.APOLLO_BLUE_LT,
         align=WD_ALIGN_PARAGRAPH.CENTER, spacing=50, after=8)

    # 표제 패널 — 옅은 틴트 바탕, 위아래 굵은 괘선
    pt = doc.add_table(rows=1, cols=1)
    table_fixed(pt, [_text_width_tw(doc)])
    cell = pt.rows[0].cells[0]
    shade_cell(cell, mt.APOLLO_PALE)
    table_grid(pt, mt.APOLLO_BLUE, sz=18, sides="h")
    table_cellmar(pt, top=90, bottom=90, left=80, right=80)
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    para(cell, COVER_TITLE1, 29, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=1.15)
    para(cell, COVER_TITLE2, 29, bold=True, color=mt.APOLLO_BLUE_MID,
         align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=2, line=1.15)
    para(cell, "◆ ————————————— ◆", 10, color=mt.APOLLO_BLUE_LT,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(cell, ENGINE_NOTE, 12.5, bold=True, color=mt.APOLLO_BLUE,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(doc, "", after=6)
    _band(doc, [mt.APOLLO_BLUE, mt.APOLLO_BLUE_MID, mt.APOLLO_BLUE_LT], after=40)

    # 하단 APOLLO 로고 — 표지에서 가장 큰 브랜드 요소
    ap = doc.add_paragraph()
    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ap.paragraph_format.space_after = Pt(6)
    try:
        ap.add_run().add_picture(mt.LOGO_APOLLO, width=Inches(COVER_APOLLO_IN))
    except Exception as e:
        print("APOLLO 로고 삽입 실패:", e)
    para(doc, "국가R&D 사업화 유망성 탐색 플랫폼", 13.5, bold=True,
         color=mt.APOLLO_BLUE_MID, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    _band(doc, [mt.APOLLO_BLUE_LT, mt.APOLLO_BLUE], h=6, after=34)
    add_disclaimer_box(doc)


# ---- 개요 · 목차 ------------------------------------------------------------
def build_intro_toc(doc, demands, n_rec, n_proj, n_sup, page_break=True):
    n_dem = len(demands)
    if page_break:                     # 구역 나누기로 이미 쪽이 넘어간 경우는 생략
        doc.add_paragraph().paragraph_format.page_break_before = True
    section_label(doc, "행사 개요", before=4)
    para(doc, mt.EVENT["명칭"], 11, bold=True, color=NAVY, after=6)
    t = doc.add_table(rows=0, cols=2)
    table_grid(t, HAIR, 4, "all"); table_cellmar(t, 40, 40, 110, 110)
    for k in mt.EVENT_ROWS:
        cells = t.add_row().cells
        shade_cell(cells[0], LABEL_BG); cell_vcenter(cells[0])
        fill_cell(cells[0], k, 9.5, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[1], mt.EVENT[k], 9, line=1.3)
    table_fixed(t, [1200, 7440])
    rows_cantsplit(t)
    para(doc, mt.EVENT["목적"], 10, color=INK, family=SERIF,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, before=8, after=2)
    para(doc, f"참가대상  {mt.EVENT['참가대상']}", 9, color=MUTED, after=1)
    para(doc, f"자세한 내용  {mt.EVENT['홈페이지']}", 9, color=MUTED, after=4)

    section_label(doc, "보고서 개요", before=16)
    overview = (
        f"본 보고서는 2026 MEDITEK 참여기업 {n_dem}개사의 기술을 대상으로, MEDITEK 공급기관 "
        f"{n_sup}곳이 수행한 국가 R&D 과제와의 의미 기반 매칭을 수행한 결과를 정리한 것이다. "
        f"기업별로 적합도가 높은 추천 과제 상위 5건(총 {n_rec}건, 중복 제외 {n_proj}개 과제)을 "
        f"선정하고, 매칭 근거와 상세 추천 근거를 함께 제시하였다. {METHOD_NOTE} {CORPUS_NOTE}"
    )
    para(doc, overview, 10.5, color=INK, family=SERIF,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.6, after=8)
    para(doc, "각 기업은 다음 순서로 구성된다.", 10.5, color=INK, family=SERIF, after=3)
    for ln in ["기업 정보  —  기업명 · 수요기술(확보 희망) · 보유기술(이미 보유) · 핵심 키워드",
               f"최종 추천 과제 {TOPN_LABEL}  —  순위 · 과제명 · 수행기관 · 공급기관 · "
               "수행년도 · 특허 · 매칭 근거",
               "추천 과제별 상세 정보표  —  과제고유번호 · 수행기간 · 표준분류 · 연구개발단계 · "
               "수행기관 · 연구수행주체 · 연구책임자 · 국가연구자번호",
               "추천 과제별 상세 매칭 근거  —  연관성 · 기술 적합성 · 추천 과제의 우수성 · "
               "유사 사례 및 실적",
               "추천 과제별 특허 실적  —  등록·출원 구분 · 특허명 · 기관 · 국가 · 번호 · 일자"]:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Twips(260)
        q.paragraph_format.space_after = Pt(3)
        style_run(q.add_run("· "), 10.5, bold=True, color=ACCENT)
        style_run(q.add_run(ln), 10, color=INK, family=SERIF)
    para(doc, "'공급기관'은 해당 과제를 수행한 기관의 기술이전 창구(산학협력단·기술지주 등)로, "
              "기술이전 협의를 시작할 접촉 지점을 뜻한다.",
         9.5, color=MUTED, family=SERIF, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         line=1.45, before=6, after=4)

    section_label(doc, "목차", before=18)
    para(doc, "기업별 시작 페이지. (Word에서 열면 자동 갱신되며, 갱신 안 되면 목차 위에서 F9)",
         8.5, color=MUTED, after=6)
    p = doc.add_paragraph()
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "2-2" \\h \\z \\u'; r._r.append(it)
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); r._r.append(sep)
    style_run(p.add_run("목차를 표시하려면 문서를 열고 F9로 필드를 업데이트하세요."), 9.5, color=MUTED)
    e = p.add_run()
    ec = OxmlElement("w:fldChar"); ec.set(qn("w:fldCharType"), "end"); e._r.append(ec)
    try:
        st = doc.settings.element
        if st.find(qn("w:updateFields")) is None:
            uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); st.append(uf)
    except Exception:
        pass


# ---- 기업 목록 · 공급기관 요약 ------------------------------------------------
def build_company_list(doc, ks, demands):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(0); h.paragraph_format.space_after = Pt(2)
    style_run(h.add_run("대상 기업 목록"), 21, bold=True, color=NAVY)
    cap = para(doc, "TARGET COMPANIES", 9, bold=True, color=ACCENT, spacing=50, after=6)
    para_border(cap, "bottom", NAVY, 14, 6)
    para(doc, f"기업 {len(ks)}개사", 10, color=MUTED, after=12)

    t = doc.add_table(rows=1, cols=3)
    table_grid(t, HAIR, 4, "all"); table_cellmar(t)
    for c, txt in zip(t.rows[0].cells, ("번호", "기업명", "기술명")):
        shade_cell(c, HEAD_BG); cell_vcenter(c)
        fill_cell(c, txt, 9.5, bold=True, color=HEAD_FG, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, k in enumerate(ks):
        v = demands[k]
        row = t.add_row().cells
        if idx % 2 == 1:
            for c in row:
                shade_cell(c, ZEBRA)
        fill_cell(row[0], k, 9.5, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(row[1], v["기업명"], 9.5, line=1.25)
        names = tech_names(v)
        fill_cell(row[2], "\n".join(names) if names else (v.get("기술분야") or ""),
                  9, line=1.25, color=INK if names else MUTED)
        for c in row:
            cell_vcenter(c)
    table_fixed(t, [700, 2000, 5940])


def build_supply_summary(doc, demands):
    doc.add_paragraph().paragraph_format.page_break_before = True
    section_label(doc, "공급기관별 추천 현황", before=0, after=5)
    para(doc, "추천 과제를 수행한 공급기관과 그 비중. 같은 과제가 여러 기업에 추천될 수 있어 "
              "추천 건수와 과제 수는 다를 수 있다.", 9, color=MUTED, after=6)
    heads = ("공급기관", "추천 건수", "과제 수", "특허 건수")
    t = doc.add_table(rows=1, cols=len(heads))
    table_grid(t, HAIR, 4, "all"); table_cellmar(t)
    for c, txt in zip(t.rows[0].cells, heads):
        shade_cell(c, HEAD_BG); cell_vcenter(c)
        fill_cell(c, txt, 9.5, bold=True, color=HEAD_FG, align=WD_ALIGN_PARAGRAPH.CENTER)
    rows = supply_summary(demands)
    for idx, (sup, a) in enumerate(rows):
        row = t.add_row().cells
        if idx % 2 == 1:
            for c in row:
                shade_cell(c, ZEBRA)
        fill_cell(row[0], sup, 9.5, line=1.25)
        fill_cell(row[1], f"{a['rec']}건", 9.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                  bold=True, color=NAVY)
        fill_cell(row[2], f"{len(a['pids'])}개", 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(row[3], f"{a['pat']}건", 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row:
            cell_vcenter(c)
    tot = t.add_row().cells
    fill_cell(tot[0], "합계", 9.5, bold=True, color=NAVY)
    fill_cell(tot[1], f"{sum(a['rec'] for _, a in rows)}건", 9.5, bold=True,
              color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(tot[2], f"{len(set().union(*(a['pids'] for _, a in rows)))}개", 9.5,
              bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(tot[3], f"{sum(a['pat'] for _, a in rows)}건", 9.5, bold=True,
              color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in tot:
        shade_cell(c, LABEL_BG); cell_vcenter(c)
    table_fixed(t, [4140, 1500, 1500, 1500])
    rows_cantsplit(t)


# ---- 기업 블록 --------------------------------------------------------------
def build_company(doc, k, dm, pidf):
    h2 = doc.add_heading(level=2)
    h2.paragraph_format.page_break_before = True
    h2.paragraph_format.space_before = Pt(0); h2.paragraph_format.space_after = Pt(5)
    badge_txt = f" No.{k} "
    _off = _text_w(badge_txt, 12) + _text_w("  ", 12)
    h2.paragraph_format.left_indent = Pt(_off)
    h2.paragraph_format.first_line_indent = Pt(-_off)
    badge = h2.add_run(badge_txt)
    style_run(badge, 12, bold=True, color="FFFFFF"); run_shade(badge, ACCENT)
    style_run(h2.add_run("  "), 12)
    style_run(h2.add_run(dm["기업명"]), 14, bold=True, color=NAVY)
    para_border(h2, "bottom", HAIR, 6, 6)

    # 기업 정보 표 — 수요기술(확보 희망)과 보유기술(이미 보유)을 구분 머리행으로 나눈다
    t = doc.add_table(rows=0, cols=2)
    table_grid(t, HAIR, 4, "all"); table_cellmar(t, 44, 44, 110, 110)

    def _row(label, val, prose=False):
        cells = t.add_row().cells
        shade_cell(cells[0], LABEL_BG); cell_vcenter(cells[0])
        fill_cell(cells[0], label, 9.5, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[1], val, 9,
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY if prose else None,
                  family=SERIF if prose else gr.SANS, line=1.28)

    def _group(gubun):
        cells = t.add_row().cells
        cell = cells[0].merge(cells[1])
        shade_cell(cell, HEAD_BG); cell_vcenter(cell)
        fill_cell(cell, GUBUN_LABEL[gubun], 9.5, bold=True, color=HEAD_FG)
        for para_ in cell.paragraphs:            # 구분 머리행 + 첫 항목을 한 페이지에
            keep_next(para_)

    _row("기업명", dm["기업명"])
    ind = industry_row(dm)
    if ind:
        _row(ind[0], ind[1])
    nm = {g: (n, gb) for g, n, gb in tech_name_rows(dm)}
    body = tech_rows(dm)
    for gubun in ("수요기술", "보유기술"):
        has_name, has_body = gubun in nm, [x for x in body if x[0] == gubun]
        if not (has_name or has_body):
            continue
        _group(gubun)
        if has_name:
            n, gb = nm[gubun]
            _row("기술명", n + (f"\n({gb})" if gb else ""))
        for _g, label, val, prose in has_body:
            _row(label, val, prose)
    if (dm.get("키워드") or "").strip():
        _row("핵심 키워드", " · ".join(x for x in dm["키워드"].split(";") if x))
    table_fixed(t, [1560, 7080])

    # 최종 추천 Top5 — 적합도 점수는 싣지 않는다(순위로만 제시)
    section_label(doc, f"최종 추천 과제  {TOPN_LABEL}", before=9, after=5)
    heads = ("순위", "과제명", "수행기관", "공급기관", "수행년도", "특허", "매칭 근거")
    t = doc.add_table(rows=1, cols=len(heads))
    table_grid(t, HAIR, 4, "all"); table_cellmar(t, 32, 32, 62, 62)
    for c, txt in zip(t.rows[0].cells, heads):
        shade_cell(c, HEAD_BG); cell_vcenter(c)
        fill_cell(c, txt, 9, bold=True, color=HEAD_FG, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, tp in enumerate(dm[TOP_KEY]):
        cells = t.add_row().cells
        if idx % 2 == 1:
            for c in cells:
                shade_cell(c, ZEBRA)
        fill_cell(cells[0], str(tp["rank"]), 11, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[1], tp["과제명"], 8.6, line=1.16)
        fill_cell(cells[2], tp.get("수행기관", ""), 8.2, align=WD_ALIGN_PARAGRAPH.CENTER,
                  line=1.14)
        fill_cell(cells[3], tp.get("공급기관", ""), 8.2, align=WD_ALIGN_PARAGRAPH.CENTER,
                  color=ACCENT, bold=True, line=1.14)
        fill_cell(cells[4], "\n".join(year_lines(tp.get("과제설명문", ""))), 8.2,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        npat = patent_count(tp)
        fill_cell(cells[5], f"{npat}건", 8.2, align=WD_ALIGN_PARAGRAPH.CENTER,
                  bold=npat > 0, color=NAVY if npat > 0 else MUTED)
        fill_cell(cells[6], gm.rename_rnd(tp.get("판단근거", "")), 8.2,
                  family=SERIF, line=1.16)
        for c in cells:
            cell_vcenter(c)
    table_fixed(t, [620, 2150, 1030, 1100, 800, 440, 2500])
    rows_cantsplit(t)

    for tp in dm[TOP_KEY]:
        gr.build_top_detail(doc, gm.neutralize(tp), pidf)


def build():
    demands, pidf = load_data()
    try:
        gr.LOGO_PATH = gr.make_top_logo()
    except Exception as e:
        print("로고 크롭 실패(로고 없이 진행):", e)
        gr.LOGO_PATH = None

    doc = Document()
    setup_styles(doc)
    set_margins(doc.sections[0])

    ks = sorted(demands, key=gm.sort_key)
    n_rec = sum(len(demands[k][TOP_KEY]) for k in ks)
    n_proj = len({t["과제고유번호"] for k in ks for t in demands[k][TOP_KEY]})
    n_sup = len({t.get("공급기관", "") for k in ks for t in demands[k][TOP_KEY]
                 if t.get("공급기관")})

    build_cover(doc)
    # 표지 전용 구역 — 헤더(작은 APOLLO 로고)를 두지 않아 하단 대형 로고와 겹치지 않게 한다.
    # 표지 전용 구역을 끊는다. add_section() 은 본문 말단의 sectPr(sentinel)을 복제해
    # 새 단락에 넣고 **sentinel 자체를 돌려준다** — 즉 여기서 얻은 객체는 다음
    # add_section() 때 마지막 구역으로 밀려난다. 그래서 구역 객체를 들고 있지 말고
    # 모든 구역을 만든 뒤 인덱스로 다시 잡는다(들고 있으면 개요 설정이 본문에 적용된다).
    doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(doc.sections[-1])
    build_intro_toc(doc, demands, n_rec, n_proj, n_sup, page_break=False)

    doc.add_section(WD_SECTION.NEW_PAGE)
    sec_cover, sec_intro, body = doc.sections[0], doc.sections[1], doc.sections[2]
    set_margins(body)
    set_pgnum_start(body, 1)
    setup_running_header(body)
    body.footer.is_linked_to_previous = False
    fp = body.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gr.add_page_field(fp)

    sec_cover.footer.is_linked_to_previous = False
    sec_cover.header.is_linked_to_previous = False            # 표지: 헤더 비움
    # 개요·목차 구역의 푸터를 **먼저** 끊고 비운다. 표지 푸터를 채운 뒤에 끊으면 그
    # 내용이 복사돼 개요 페이지까지 버전 표기가 따라붙는다(실측).
    sec_intro.footer.is_linked_to_previous = False
    for _p in sec_intro.footer.paragraphs:
        for _r in list(_p.runs):
            _r._element.getparent().remove(_r._element)
    sec_intro.header.is_linked_to_previous = False
    for _p in sec_intro.header.paragraphs:
        for _r in list(_p.runs):
            _r._element.getparent().remove(_r._element)
    setup_cover_header(sec_intro)                             # 개요·목차: 로고만
    # 표지 우측 하단 버전 표기 — 구역 푸터에 넣어 페이지 하단에 고정한다
    cfp = sec_cover.footer.paragraphs[0]
    cfp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cfp.paragraph_format.space_before = Pt(0)
    style_run(cfp.add_run(COVER_VERSION), 7.5, color=MUTED)

    build_company_list(doc, ks, demands)
    build_supply_summary(doc, demands)
    for k in ks:
        build_company(doc, k, demands[k], pidf)

    doc.save(OUT)
    print("saved:", OUT)
    print(f"  기업 {len(ks)}개사 · 추천 {n_rec}건 · 중복제외 과제 {n_proj}개 · 공급기관 {n_sup}곳")
    return OUT


if __name__ == "__main__":
    build()
