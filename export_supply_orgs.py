# -*- coding: utf-8 -*-
"""공급기관 → 과제수행기관명 대응표를 엑셀로 저장(+ 코퍼스 대조 감사).

meditek_supply_orgs.SUPPLY_ORGS 를 코퍼스(project_match_data_260612.pkl)와 맞춰 보고
과제 건수·매칭 대상 건수를 붙여 3개 시트로 낸다.

  대응표      공급기관 1행 = 과제수행기관 1개(원본 기관명·부서·범위 규칙 포함)
  공급기관요약 공급기관별 기관 수·과제 합계·매칭 대상 과제 수
  제외검토    PROBE 에 걸렸지만 포함하지 않은 기관명과 제외 사유(감사용)

'매칭 대상'은 실제 매칭 코퍼스 조건과 같다: 제출년도>=YEAR_MIN ∧ 특허 성과 1건 이상.

사용: python export_supply_orgs.py [--out 2026_MEDITEK_공급기관_과제수행기관.xlsx]
"""
import argparse
import collections
import pickle
import re

import numpy as np
import pandas as pd

import compa_match as cm
import match_meditek_supply as ms
import meditek_supply_orgs as so
from rematch_filtered import EMB_FILE, YEAR_MIN

SRC_XLSX = "2026_MEDITEK_260820.xlsx"
OUT = "2026_MEDITEK_공급기관_과제수행기관.xlsx"

# 제외 사유 — PROBE 에 걸린 이름을 사람이 확인해 분류한 근거(감사 흔적)
EXCLUDE_REASON = [
    (r"^대구경북첨단의료산업진흥재단|^\(재단\)대구경북", "별개 재단(대구경북)"),
    (r"^(대구|부산|인천|목포|수원)가톨릭|^가톨릭관동|^가톨릭상지|국제성모병원", "별개 대학·부속병원"),
    (r"^고려대학교$|^고려대$|^고려대\(조치원\)|세종캠퍼스|기환경연구소|기술지주|산학협력단",
     "고려대 병원 계열 아님(요청에 따라 제외)"),
    (r"^한양대학교$|^한양대$|한양대학교병원|한양대기술지주|한양대산학협력단",
     "ERICA 캠퍼스 아님(요청에 따라 제외)"),
    (r"^울산과학대", "별개 기관(전문대) — UNIST 아님"),
    (r"^\(학교\)수원인제학원$", "별개 학교법인(수원인제학원)"),
    (r"^숭실사이버대", "별개 대학(원격대학)"),
    (r"^안동과학대|^경도대학교$|^전남도립|^강원도립|^충북도립|^충남도립", "별개 대학"),
    (r"^동부산대|^양산대학교$", "별개 대학"),
    (r"산학협력단", "본교 산학협력단(범위 외)"),
    (r"병원|의료원", "타 기관 병원·의료원"),
    (r"대학교|대학$|대학교$", "타 대학"),
]


def reason_for(name):
    for pat, why in EXCLUDE_REASON:
        if re.search(pat, name):
            return why
    return "무관 기관·기업(이름만 유사)"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=SRC_XLSX)
    ap.add_argument("--out", default=OUT)
    ap.add_argument("--year-min", type=int, default=YEAR_MIN)
    a = ap.parse_args()

    print("· 과제 메타 로드…")
    with open(cm.PROJECT_META, "rb") as f:
        pm = pickle.load(f)
    org_of = {p: str(v.get("과제수행기관명", "")) for p, v in pm.items()}
    cnt_all = collections.Counter(org_of.values())

    missing, unclassified = so.audit(cnt_all)
    if missing:
        raise SystemExit(f"[중단] 코퍼스에 없는 포함대상: {missing}")
    print(f"  고유 과제수행기관명 {len(cnt_all)}개 · 포함대상 전부 코퍼스에 존재")

    print("· 매칭 대상(제출년도·특허) 산정…")
    pdf = pd.read_pickle(EMB_FILE)
    pid = pdf["과제고유번호"].astype(str).values
    year = pd.to_numeric(pdf["제출년도"], errors="coerce").values
    pcnt = ms.patent_counts()
    org_arr = np.array([org_of.get(p, "") for p in pid])
    ok = (year >= a.year_min) & np.array([pcnt.get(p, 0) >= 1 for p in pid])
    cnt_ok = collections.Counter(org_arr[ok])
    del pdf

    sup_df = pd.read_excel(a.src, sheet_name="공급기관 리스트")
    meta = {str(r["기관명"]).strip(): (r.get("No"), str(r.get("기관유형", "")).strip(),
                                      str(r.get("부서", "")).strip())
            for r in sup_df.to_dict("records")}

    rows, summ, excl = [], [], []
    for sup, (memo, probe, orgs) in so.SUPPLY_ORGS.items():
        no, typ, dept = meta.get(sup, ("", "", ""))
        if not orgs:
            rows.append({"No": no, "공급기관명(원본)": sup, "기관유형": typ, "부서": dept,
                         "범위 규칙": memo, "과제수행기관명": "(해당 없음)",
                         "과제 건수": 0, "매칭 대상 과제": 0})
        for o in sorted(orgs, key=lambda x: -cnt_all.get(x, 0)):
            rows.append({"No": no, "공급기관명(원본)": sup, "기관유형": typ, "부서": dept,
                         "범위 규칙": memo, "과제수행기관명": o,
                         "과제 건수": cnt_all.get(o, 0), "매칭 대상 과제": cnt_ok.get(o, 0)})
        summ.append({"No": no, "공급기관명(원본)": sup, "기관유형": typ, "부서": dept,
                     "범위 규칙": memo, "과제수행기관 수": len(orgs),
                     "과제 건수 합계": sum(cnt_all.get(o, 0) for o in orgs),
                     "매칭 대상 과제 합계": sum(cnt_ok.get(o, 0) for o in orgs)})
        for o in unclassified.get(sup, []):
            excl.append({"공급기관명(원본)": sup, "제외한 과제수행기관명": o,
                         "과제 건수": cnt_all.get(o, 0), "제외 사유": reason_for(o)})

    df = pd.DataFrame(rows)
    ds = pd.DataFrame(summ).sort_values("매칭 대상 과제 합계", ascending=False)
    de = pd.DataFrame(excl).sort_values(["공급기관명(원본)", "과제 건수"],
                                        ascending=[True, False])
    tot = {"공급기관명(원본)": "합계", "과제수행기관 수": ds["과제수행기관 수"].sum(),
           "과제 건수 합계": ds["과제 건수 합계"].sum(),
           "매칭 대상 과제 합계": ds["매칭 대상 과제 합계"].sum()}
    ds = pd.concat([ds, pd.DataFrame([tot])], ignore_index=True)

    W = {"No": 6, "공급기관명(원본)": 26, "기관유형": 9, "부서": 20, "범위 규칙": 62,
         "과제수행기관명": 40, "과제 건수": 10, "매칭 대상 과제": 13,
         "과제수행기관 수": 14, "과제 건수 합계": 13, "매칭 대상 과제 합계": 16,
         "제외한 과제수행기관명": 40, "제외 사유": 34}
    WRAP = {"범위 규칙", "과제수행기관명", "제외한 과제수행기관명", "부서", "제외 사유"}
    with pd.ExcelWriter(a.out, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="대응표", index=False)
        ds.to_excel(xw, sheet_name="공급기관요약", index=False)
        de.to_excel(xw, sheet_name="제외검토", index=False)
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = load_workbook(a.out)
    for name, d in (("대응표", df), ("공급기관요약", ds), ("제외검토", de)):
        ws = wb[name]
        for j, col in enumerate(d.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = W.get(col, 14)
            if col in WRAP:
                for i in range(2, ws.max_row + 1):
                    ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9E1F2")
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
    wb.save(a.out)

    print(f"\n저장: {a.out}")
    print(f"  대응표 {len(df)}행 · 공급기관 {len(so.SUPPLY_ORGS)}곳 · "
          f"과제수행기관 {ds['과제수행기관 수'].iloc[-1]}개")
    print(f"  과제 합계 {ds['과제 건수 합계'].iloc[-1]:,}건 → "
          f"매칭 대상(제출년도>={a.year_min} ∧ 특허1건+) {ds['매칭 대상 과제 합계'].iloc[-1]:,}건")
    print(f"  제외검토 {len(de)}행")
    print()
    print(ds.to_string(index=False))


if __name__ == "__main__":
    main()
