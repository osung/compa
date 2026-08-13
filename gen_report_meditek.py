# -*- coding: utf-8 -*-
"""2026 MEDITEK 수요기술 매칭 보고서(docx) 생성.

gen_report.py 의 MEDITEK 판. 서식·수요 블록·과제 상세 블록은 gen_report 를 그대로
재사용하고, 아래만 다르게 구성한다.
  - 6T 분야 분류 없음 → 분야별 장(章) 구분 없이 수요기술을 일련으로 수록
  - 표지/개요/러닝 문구를 MEDITEK 용으로 교체(COMPA 매칭데이 고정 문구 제거)

입력:
  MEDITEK_통합best.json          최종 추천 Top5 + 근거(모델 생성)  ← build_meditek_best.py
  $COMPA_SCRATCH/pid_fields.json 연구개발단계/연구수행주체         ← _proj_meta_prep.py
  $COMPA_SCRATCH/pid_patents.json 특허 실적(선택)                 ← _patent_prep.py
  COMPA_보고서_작성규칙.md         서식 규칙

사용: COMPA_SCRATCH=<scratch> python gen_report_meditek.py
"""
import glob
import json
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

import gen_report as gr
from gen_report import (ACCENT, HAIR, HEAD_BG, HEAD_FG, INK, MUTED, NAVY,
                        SERIF, ZEBRA, add_disclaimer_box, build_demand, cell_vcenter,
                        fill_cell, para, para_border, section_label, set_margins,
                        set_pgnum_start, setup_cover_header, setup_running_header,
                        setup_styles, shade_cell, style_run, table_cellmar,
                        table_fixed, table_grid)

HERE = os.path.dirname(os.path.abspath(__file__))
BEST_JSON = os.environ.get("MEDITEK_REPORT_JSON", os.path.join(HERE, "MEDITEK_통합best.json"))
OUT_BASE = "MEDITEK_수요기술_매칭_보고서"
PUBLISH_DATE = os.environ.get("MEDITEK_PUBLISH_DATE", "2026. 8. 11.")

# 표지·개요 문구(MEDITEK)
COVER_EYEBROW = "TECHNOLOGY  DEMAND  ×  NATIONAL  R&D  MATCHING"
COVER_TITLE1 = "2026 MEDITEK"
COVER_TITLE2 = "수요기술 매칭 보고서"
COVER_SUB = "참여기업 수요기술 × 국가 R&D 과제, 의미 기반 매칭 결과"
ENGINE_NOTE = "생성  APOLLO AI 매칭 엔진"

# 매칭 대상 코퍼스(rematch_filtered.build_filtered_corpus 와 일치)
CORPUS_NOTE = ("매칭 대상 과제는 연구수행주체가 대학·출연연구소·국공립연구소·정부부처인 "
               "최근 5년(2020년 이후 제출) 국가 R&D 과제로 한정하였다.")


def next_out_path():
    vers = [int(m.group(1)) for p in glob.glob(os.path.join(HERE, OUT_BASE + "_v*.docx"))
            if (m := re.search(r"_v(\d+)\.docx$", p))]
    return os.path.join(HERE, f"{OUT_BASE}_v{(max(vers) + 1) if vers else 1}.docx")


OUT = os.environ.get("MEDITEK_REPORT_OUT", next_out_path())


# 과제 상세 정보표에서 제외할 항목 — 원본 데이터가 없어 '-' 만 찍히므로 싣지 않는다.
OMIT_INFO_FIELDS = {"연구책임자", "국가연구자번호"}
gr.OMIT_INFO_FIELDS = OMIT_INFO_FIELDS

# 정보표에 실제로 실리는 필드만 점검 → (pid_fields 키, 표시 라벨)
REQUIRED_PID_FIELDS = [(k, lab) for k, lab in
                       [("연구개발단계", "연구개발단계"), ("연구수행주체", "연구수행주체"),
                        ("연구책임자명", "연구책임자"), ("국가연구자번호", "국가연구자번호")]
                       if lab not in OMIT_INFO_FIELDS]


def preflight(demands, pidf):
    """정보표·특허가 통째로 비는 사고를 막기 위한 사전 점검.

    보고서는 값이 없으면 조용히 '-'로 찍히므로, 생성 전에 커버리지를 출력하고
    0%(전량 누락)인 항목은 경고한다. 부분 누락은 실제 데이터 공백일 수 있어 알리기만 한다.
    """
    pids = sorted({t["과제고유번호"] for v in demands.values() for t in v["top5"]})
    n = len(pids)
    # pidf 는 렌더러가 pid 로 바로 조회하는 dict 그대로여야 한다(감싼 구조면 전량 '-' 가 된다).
    fields = pidf
    if not (set(pids) & set(fields)):
        raise SystemExit("[preflight] pid_fields 에서 과제를 하나도 찾지 못했습니다. "
                         "구조가 {pid: {...}} 인지 확인하세요(‘fields’ 로 감싸져 있으면 벗겨야 함).")
    print(f"[preflight] 과제 {n}건 · 수요 {len(demands)}건")
    missing = []
    for key, label in REQUIRED_PID_FIELDS:
        c = sum(1 for p in pids if str(fields.get(p, {}).get(key, "") or "").strip())
        mark = "OK " if c == n else ("!! " if c == 0 else "~  ")
        print(f"  {mark}{label:<12} {c}/{n}")
        if c == 0:
            missing.append(label)

    # 설명문에서 뽑는 항목은 pid_fields 로 세지 못하므로 실제 표시값을 그대로 계산해 본다.
    tops = [t for v in demands.values() for t in v["top5"]]
    derived = {
        "과제수행기간": lambda t: gr.fmt_period(t.get("과제설명문", "")),
        "과학기술표준분류(중)": lambda t: (gr.extract_class(t.get("과제설명문", ""))
                                 or fields.get(t["과제고유번호"], {}).get("표준분류중")),
        "과제수행기관": lambda t: t.get("수행기관", ""),
    }
    for label, fn in derived.items():
        if label in OMIT_INFO_FIELDS:
            continue
        c = sum(1 for t in tops if str(fn(t) or "").strip())
        mark = "OK " if c == len(tops) else ("!! " if c == 0 else "~  ")
        print(f"  {mark}{label:<12} {c}/{len(tops)} (Top5 기준)")
        if c == 0:
            missing.append(label)
    pats = sum(1 for p in pids if gr.PATENTS.get(p))
    print(f"  {'OK ' if pats else '!! '}특허 실적       {pats}/{n} 과제 보유"
          f" (총 {sum(len(gr.PATENTS.get(p, [])) for p in pids)}건)")
    if not pats:
        missing.append("특허 실적")
    if missing:
        print(f"  ※ 전량 누락: {', '.join(missing)} — 원본 데이터/전처리 스크립트 확인 필요")
    return missing


def load_data():
    """6T(demand_field.json) 없이 수요 JSON + 과제 메타만 로드."""
    with open(BEST_JSON, encoding="utf-8") as f:
        demands = json.load(f)
    pf = os.path.join(gr.SCRATCH, "pid_fields.json")
    if not os.path.exists(pf):
        raise SystemExit(f"없는 파일: {pf} — 먼저 `python _proj_meta_prep.py` 실행 필요")
    with open(pf, encoding="utf-8") as f:
        pidf = json.load(f)
    # _proj_meta_prep.py 는 {"name2pid":…, "fields": {pid: {...}}} 로 저장한다.
    # 상세 정보표는 pid 로 바로 조회하므로 fields 를 벗겨서 넘겨야 값이 들어간다.
    pidf = pidf.get("fields", pidf)
    if not gr.PATENTS:
        print(f"! 특허 데이터 없음({gr.SCRATCH}/pid_patents.json)\n"
              f"  COMPA_REPORT_JSON={BEST_JSON} 로 `python _patent_prep.py` 를 먼저 실행하세요.")
    preflight(demands, pidf)
    return demands, pidf


# ---- 표지 -------------------------------------------------------------------
def build_cover(doc, n_dem, n_rec, n_proj):
    para(doc, "", after=54)
    para(doc, COVER_EYEBROW, 10.5, bold=True, color=ACCENT,
         align=WD_ALIGN_PARAGRAPH.CENTER, spacing=60, after=10)
    para(doc, COVER_TITLE1, 27, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, COVER_TITLE2, 27, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    rule = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    para_border(rule, "bottom", NAVY, 18, 2)
    para(doc, COVER_SUB, 12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=44)

    t = doc.add_table(rows=2, cols=3)
    t.alignment = 1
    for j, (lab, val) in enumerate([("대상 수요기술", f"{n_dem}건"),
                                    ("추천 과제", f"{n_rec}건"),
                                    ("중복 제외 과제", f"{n_proj}개")]):
        fill_cell(t.rows[0].cells[j], val, 20, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(t.rows[1].cells[j], lab, 9.5, color=MUTED,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    table_fixed(t, [2400, 2400, 2400]); table_cellmar(t, 30, 30, 60, 60)
    para(doc, "", after=40)
    para(doc, f"발행일  {PUBLISH_DATE}      {ENGINE_NOTE}",
         10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    para(doc, "", after=30)
    add_disclaimer_box(doc)


# ---- 개요 · 목차 ------------------------------------------------------------
def build_intro_toc(doc, n_dem, n_rec, n_proj):
    doc.add_paragraph().paragraph_format.page_break_before = True
    section_label(doc, "개요", before=4)
    overview = (
        f"본 보고서는 2026 MEDITEK 참여기업이 제출한 수요기술 {n_dem}건을 대상으로, "
        f"국가 R&D 과제 데이터베이스와의 의미 기반 매칭을 수행한 결과를 정리한 것이다. "
        f"각 수요기술에 대해 적합도가 높은 추천 과제 상위 5건(총 {n_rec}건, "
        f"중복 제외 {n_proj}개 과제)을 선정하고, 매칭 근거와 상세 추천 근거를 함께 제시하였다. "
        f"{CORPUS_NOTE}"
    )
    para(doc, overview, 10.5, color=INK, family=SERIF,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.6, after=8)
    para(doc, "각 수요기술은 다음 순서로 구성된다.", 10.5, color=INK, family=SERIF, after=3)
    for ln in ["수요 정보  —  기업명 · 수요기술 내용",
               "최종 추천 과제 Top 5  —  순위 · 과제명 · 수행기관 · 과제수행년도 · 매칭 근거",
               "추천 과제별 상세 정보표 및 상세 매칭 근거  —  연관성 · 수요기술 사양 적합성 · "
               "추천 과제의 우수성 · 유사 사례 및 실적"]:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Twips(260)
        q.paragraph_format.space_after = Pt(3)
        style_run(q.add_run("· "), 10.5, bold=True, color=ACCENT)
        style_run(q.add_run(ln), 10, color=INK, family=SERIF)

    section_label(doc, "목차", before=20)
    para(doc, "수요기술별 시작 페이지. (Word에서 열면 자동 갱신되며, 갱신 안 되면 목차 위에서 F9)",
         8.5, color=MUTED, after=6)
    p = doc.add_paragraph()
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "2-2" \\h \\z \\u'; r._r.append(it)   # 장 없음 → Heading 2 만
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); r._r.append(sep)
    style_run(p.add_run("목차를 표시하려면 문서를 열고 F9로 필드를 업데이트하세요."), 9.5, color=MUTED)
    e = p.add_run()
    ec = OxmlElement("w:fldChar"); ec.set(qn("w:fldCharType"), "end"); e._r.append(ec)
    try:
        st = doc.settings.element
        if st.find(qn("w:updateFields")) is None:
            uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); st.append(uf)
    except Exception:
        pass


# ---- 수요기술 목록(장 오프너 대체) ------------------------------------------
def build_demand_list(doc, ks, demands):
    """분야별 장 대신, 본문 첫 페이지에 전체 수요기술 목록을 싣는다."""
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(0); h.paragraph_format.space_after = Pt(2)
    style_run(h.add_run("수요기술 목록"), 21, bold=True, color=NAVY)
    cap = para(doc, "TECHNOLOGY DEMANDS", 9, bold=True, color=ACCENT, spacing=50, after=6)
    para_border(cap, "bottom", NAVY, 14, 6)
    para(doc, f"수요기술 {len(ks)}건", 10, color=MUTED, after=12)

    t = doc.add_table(rows=1, cols=3)
    table_grid(t, HAIR, 4, "all"); table_cellmar(t)
    for c, txt in zip(t.rows[0].cells, ("번호", "수요기술명", "기업명")):
        shade_cell(c, HEAD_BG); cell_vcenter(c)
        fill_cell(c, txt, 9.5, bold=True, color=HEAD_FG, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, k in enumerate(ks):
        row = t.add_row().cells
        if idx % 2 == 1:
            for c in row:
                shade_cell(c, ZEBRA)
        fill_cell(row[0], k, 9.5, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(row[1], demands[k]["수요기술명"], 9.5, line=1.25)
        fill_cell(row[2], demands[k]["기업명"], 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row:
            cell_vcenter(c)
    table_fixed(t, [760, 5080, 2800])


def build():
    demands, pidf = load_data()
    try:
        gr.LOGO_PATH = gr.make_top_logo()
    except Exception as e:
        print("로고 크롭 실패(로고 없이 진행):", e)
        gr.LOGO_PATH = None

    doc = Document()
    setup_styles(doc)
    set_margins(doc.sections[0])

    ks = sorted(demands, key=int)
    n_rec = sum(len(demands[k]["top5"]) for k in ks)
    n_proj = len({t["과제고유번호"] for k in ks for t in demands[k]["top5"]})

    build_cover(doc, len(ks), n_rec, n_proj)
    build_intro_toc(doc, len(ks), n_rec, n_proj)

    # ===== 본문 구역 =====
    doc.add_section(WD_SECTION.NEW_PAGE)
    body = doc.sections[-1]
    set_margins(body)
    set_pgnum_start(body, 1)
    setup_running_header(body)
    body.footer.is_linked_to_previous = False
    fp = body.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gr.add_page_field(fp)
    doc.sections[0].footer.is_linked_to_previous = False
    setup_cover_header(doc.sections[0])

    build_demand_list(doc, ks, demands)
    for k in ks:
        # 제목=기업명 → Word 목차·러닝헤더가 기업명으로 표시된다
        # (수요기술명은 원본에 없어 내용 첫 줄을 규칙 추출한 값이라 식별성이 떨어짐)
        build_demand(doc, k, demands[k], pidf, title=demands[k]["기업명"])

    doc.save(OUT)
    print("saved:", OUT)
    print(f"  수요 {len(ks)}건 · 추천 {n_rec}건 · 중복제외 과제 {n_proj}개")
    return OUT


if __name__ == "__main__":
    build()
