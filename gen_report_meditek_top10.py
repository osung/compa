# -*- coding: utf-8 -*-
"""2026 MEDITEK 기업 기술 × 국가 R&D 과제 매칭 보고서(docx) — Top10 판.

gen_report_meditek.py 의 확장. 서식·과제 상세 블록은 gen_report 를 그대로 재사용하고,
아래가 다르다.
  - 매칭 단위가 수요기술이 아니라 기업 → 기업 정보 블록에 기술 정보를 한데 싣는다
  - 추천 과제 Top5 → Top10, 표에 특허건수를 함께 싣는다(적합도 점수는 싣지 않는다)

※ 보고서에는 매칭 기준(수요기술/기보유기술 중 무엇을 근거로 매칭했는지)과 적합도 점수를
  노출하지 않는다. 순위(Top N)로만 제시한다.
  기술 정보는 라벨 없이 '기술명/기술 내용'으로만 싣고, 근거 섹션 제목도 통일한다
  (SECTION_RENAME — 이전에 생성된 근거 텍스트의 섹션 제목까지 렌더링 단계에서 바꾼다).

입력:
  MEDITEK_TOP10_보고서.json        Top10 + 근거              ← explain_meditek_top10.py
  $COMPA_SCRATCH/pid_fields.json  연구개발단계/연구수행주체    ← _proj_meta_prep.py
  $COMPA_SCRATCH/pid_patents.json 특허 실적                  ← _patent_prep.py

사용: COMPA_SCRATCH=<scratch> python gen_report_meditek_top10.py
"""
import glob
import json
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

import gen_report as gr
from gen_report import (ACCENT, HAIR, HEAD_BG, HEAD_FG, INK, LABEL_BG, MUTED, NAVY,
                        SERIF, ZEBRA, add_disclaimer_box, cell_indent, cell_vcenter,
                        fill_cell, para, para_border, run_shade, section_label,
                        set_margins, set_pgnum_start, setup_cover_header,
                        setup_running_header, setup_styles, shade_cell, style_run,
                        table_cellmar, table_fixed, table_grid, rows_cantsplit,
                        year_lines, _text_w)

HERE = os.path.dirname(os.path.abspath(__file__))
BEST_JSON = os.environ.get("MEDITEK_TOP10_JSON",
                           os.path.join(HERE, "MEDITEK_TOP10_보고서.json"))
OUT_BASE = "MEDITEK_기업기술_국가RnD_매칭보고서"
PUBLISH_DATE = os.environ.get("MEDITEK_PUBLISH_DATE", "2026. 8. 12.")

COVER_EYEBROW = "COMPANY  TECHNOLOGY  ×  NATIONAL  R&D  MATCHING"
COVER_TITLE1 = "2026 MEDITEK"
COVER_TITLE2 = "국가 R&D 과제 매칭 보고서"
COVER_SUB = "참여기업 기술 × 국가 R&D 과제, 의미 기반 매칭 결과"
ENGINE_NOTE = "생성  APOLLO AI 매칭 엔진"

CORPUS_NOTE = ("매칭 대상 과제는 연구수행주체가 대학·출연연구소·국공립연구소·정부부처인 "
               "최근 5년(2020년 이후 제출) 국가 R&D 과제로 한정하였다.")
METHOD_NOTE = ("매칭은 기업이 제출한 기술 정보에서 핵심 기술 키워드를 추출하고, 그 키워드로 "
               "과제 임베딩과의 의미 유사도가 높은 후보를 선별한 뒤, 각 후보가 해당 기업의 "
               "기술과 실질적으로 부합·기여하는 정도를 평가하는 방식으로 수행하였다. "
               "최종 순위는 이 적합도에 과제의 연구성과(특허 성과를 중심으로 논문 성과)를 "
               "함께 반영해 결정하였다.")

# 근거 텍스트의 섹션 제목 통일 — 매칭 기준이 드러나는 제목을 렌더링 단계에서 바꾼다
# (이미 생성된 추천근거_상세 에는 옛 제목이 남아 있어도 보고서에는 노출되지 않는다).
SECTION_RENAME = {"수요 충족 가능성": "기술 적합성",
                  "기보유기술 보강 가능성": "기술 적합성",
                  "수요기술 사양 적합성": "기술 적합성"}


_PUBLIC_RND = re.compile(r"공공\s*R&D")     # 명칭 통일: 공공 R&D → 국가 R&D


def rename_rnd(s):
    return _PUBLIC_RND.sub(lambda m: m.group(0).replace("공공", "국가"), str(s or ""))


def neutralize(tp):
    """과제 상세 dict 사본 — 섹션 제목 통일 + '공공 R&D' 표기를 '국가 R&D' 로."""
    txt = rename_rnd(tp.get("추천근거_상세", ""))
    for old, new in SECTION_RENAME.items():
        txt = txt.replace(f"[{old}]", f"[{new}]")
    return dict(tp, 추천근거_상세=txt, 판단근거=rename_rnd(tp.get("판단근거", "")))


def patent_count(tp):
    """표에 싣는 특허 건수 — 상세 페이지의 특허 실적 목록과 같은 소스에서 센다."""
    return len(gr.PATENTS.get(str(tp["과제고유번호"]), []))


def tech_names(dm):
    """기업 기술명 — 매칭 기준 구분 없이 확보된 기술명을 그대로 나열."""
    return [x for x in (dm.get("수요기술명", ""), dm.get("기보유기술명", "")) if x.strip()]


def tech_bodies(dm):
    """기업 기술 내용 — 라벨 없이 이어붙인다(어느 쪽이 매칭 기준인지 드러내지 않음).

    제출 내용이 한 줄뿐이면 기술명과 같은 문장이 되므로(기술명은 내용 첫 줄에서 도출)
    그런 본문은 싣지 않는다 — 같은 문장이 두 행에 반복되는 것을 막는다.
    """
    def norm(s):
        return re.sub(r"\s+", "", s or "")

    names = {norm(x) for x in tech_names(dm)}
    out = []
    for x in (dm.get("수요기술 내용", ""), dm.get("기보유기술 내용", "")):
        x = (x or "").strip()
        if x and norm(x) not in names:
            out.append(x)
    return out

# 과제 상세 정보표는 8개 항목 전부 싣는다. 연구책임자·국가연구자번호는
# ntis_nrsno_260610.pkl 에서 확보되므로(_proj_meta_prep.py) 더 이상 제외하지 않는다.
gr.OMIT_INFO_FIELDS = set()


def next_out_path():
    vers = [int(m.group(1)) for p in glob.glob(os.path.join(HERE, OUT_BASE + "_v*.docx"))
            if (m := re.search(r"_v(\d+)\.docx$", p))]
    return os.path.join(HERE, f"{OUT_BASE}_v{(max(vers) + 1) if vers else 1}.docx")


OUT = os.environ.get("MEDITEK_TOP10_OUT", next_out_path())


def sort_key(k):
    """수요 번호 순 → 수요 행이 없는 기업(관리번호 순)."""
    if re.fullmatch(r"\d+", k):
        return (0, int(k))
    m = re.search(r"(\d+)$", k)
    return (1, int(m.group(1)) if m else 0)


def load_data():
    with open(BEST_JSON, encoding="utf-8") as f:
        demands = json.load(f)
    pf = os.path.join(gr.SCRATCH, "pid_fields.json")
    if not os.path.exists(pf):
        raise SystemExit(f"없는 파일: {pf} — 먼저 `python _proj_meta_prep.py` 실행 필요")
    with open(pf, encoding="utf-8") as f:
        pidf = json.load(f)
    pidf = pidf.get("fields", pidf)         # {pid: {...}} 로 벗겨서 넘겨야 값이 들어간다
    if not gr.PATENTS:
        print(f"! 특허 데이터 없음({gr.SCRATCH}/pid_patents.json)\n"
              f"  COMPA_REPORT_JSON={BEST_JSON} 로 `python _patent_prep.py` 실행 필요")
    preflight(demands, pidf)
    return demands, pidf


def preflight(demands, pidf):
    """정보표·근거·특허가 통째로 비는 사고 방지(보고서는 값이 없으면 조용히 '-' 로 찍힌다)."""
    tops = [t for v in demands.values() for t in v["top10"]]
    pids = sorted({t["과제고유번호"] for t in tops})
    # 유망성 점수 비노출 점검 — 입력 JSON·특허 데이터 어디에도 남아 있으면 안 된다
    leak = [k for t in tops for k in t if "유망" in k]
    leak += [k for lst in gr.PATENTS.values() for x in lst for k in x if "유망" in k]
    if leak:
        raise SystemExit(f"[preflight] 유망성 정보가 보고서 입력에 남아 있습니다: {set(leak)}")
    if not (set(pids) & set(pidf)):
        raise SystemExit("[preflight] pid_fields 에서 과제를 하나도 찾지 못했습니다 — "
                         "구조가 {pid: {...}} 인지 확인(‘fields’ 로 감싸져 있으면 벗겨야 함).")
    print(f"[preflight] 기업 {len(demands)}건 · 추천 {len(tops)}건 · 과제 {len(pids)}개")
    for key, label in (("연구개발단계", "연구개발단계"), ("연구수행주체", "연구수행주체"),
                       ("연구책임자명", "연구책임자"), ("국가연구자번호", "국가연구자번호")):
        c = sum(1 for p in pids if str(pidf.get(p, {}).get(key, "") or "").strip())
        print(f"  {'OK ' if c == len(pids) else ('!! ' if c == 0 else '~  ')}{label:<12} "
              f"{c}/{len(pids)}")
    for label, fn in (("과제수행기간", lambda t: gr.fmt_period(t.get("과제설명문", ""))),
                      ("과학기술표준분류(중)",
                       lambda t: (gr.extract_class(t.get("과제설명문", ""))
                                  or pidf.get(t["과제고유번호"], {}).get("표준분류중"))),
                      ("과제수행기관", lambda t: t.get("수행기관", "")),
                      ("매칭 근거", lambda t: t.get("판단근거", "")),
                      ("상세 매칭 근거", lambda t: t.get("추천근거_상세", ""))):
        c = sum(1 for t in tops if str(fn(t) or "").strip())
        print(f"  {'OK ' if c == len(tops) else ('!! ' if c == 0 else '~  ')}{label:<12} "
              f"{c}/{len(tops)}")
    pats = sum(1 for p in pids if gr.PATENTS.get(p))
    print(f"  {'OK ' if pats else '!! '}특허 실적       {pats}/{len(pids)} 과제 보유"
          f" (총 {sum(len(gr.PATENTS.get(p, [])) for p in pids)}건)")


# ---- 표지 -------------------------------------------------------------------
def build_cover(doc, n_dem, n_rec, n_proj):
    para(doc, "", after=54)
    para(doc, COVER_EYEBROW, 9.5, bold=True, color=ACCENT,
         align=WD_ALIGN_PARAGRAPH.CENTER, spacing=50, after=10)
    para(doc, COVER_TITLE1, 27, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, COVER_TITLE2, 27, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    rule = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    para_border(rule, "bottom", NAVY, 18, 2)
    para(doc, COVER_SUB, 12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=44)

    t = doc.add_table(rows=2, cols=3)
    t.alignment = 1
    for j, (lab, val) in enumerate([("대상 기업", f"{n_dem}개사"),
                                    ("추천 과제", f"{n_rec}건"),
                                    ("중복 제외 과제", f"{n_proj}개")]):
        fill_cell(t.rows[0].cells[j], val, 20, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(t.rows[1].cells[j], lab, 9.5, color=MUTED,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    table_fixed(t, [2400, 2400, 2400]); table_cellmar(t, 30, 30, 60, 60)
    para(doc, "", after=40)
    para(doc, f"발행일  {PUBLISH_DATE}      {ENGINE_NOTE}",
         10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    para(doc, "", after=30)
    add_disclaimer_box(doc)


# ---- 개요 · 목차 ------------------------------------------------------------
def build_intro_toc(doc, demands, n_rec, n_proj):
    n_dem = len(demands)
    doc.add_paragraph().paragraph_format.page_break_before = True
    section_label(doc, "개요", before=4)
    overview = (
        f"본 보고서는 2026 MEDITEK 참여기업 {n_dem}개사의 기술을 대상으로, 국가 R&D 과제 "
        f"데이터베이스와의 의미 기반 매칭을 수행한 결과를 정리한 것이다. "
        f"기업별로 적합도가 높은 추천 과제 상위 10건(총 {n_rec}건, 중복 제외 {n_proj}개 과제)을 "
        f"선정하고, 매칭 근거와 상세 추천 근거를 함께 제시하였다. {METHOD_NOTE} {CORPUS_NOTE}"
    )
    para(doc, overview, 10.5, color=INK, family=SERIF,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.6, after=8)
    para(doc, "각 기업은 다음 순서로 구성된다.", 10.5, color=INK, family=SERIF, after=3)
    for ln in ["기업 정보  —  기업명 · 기술명 · 기술 내용 · 핵심 키워드",
               "최종 추천 과제 Top 10  —  순위 · 과제명 · 수행기관 · 수행년도 · 특허 · 매칭 근거",
               "추천 과제별 상세 정보표  —  과제고유번호 · 수행기간 · 표준분류 · 연구개발단계 · "
               "수행기관 · 연구수행주체 · 연구책임자 · 국가연구자번호",
               "추천 과제별 상세 매칭 근거  —  연관성 · 기술 적합성 · 추천 과제의 우수성 · "
               "유사 사례 및 실적",
               "추천 과제별 특허 실적  —  등록·출원 구분 · 특허명 · 기관 · 국가 · 번호 · 일자"]:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Twips(260)
        q.paragraph_format.space_after = Pt(3)
        style_run(q.add_run("· "), 10.5, bold=True, color=ACCENT)
        style_run(q.add_run(ln), 10, color=INK, family=SERIF)

    section_label(doc, "목차", before=18)
    para(doc, "기업별 시작 페이지. (Word에서 열면 자동 갱신되며, 갱신 안 되면 목차 위에서 F9)",
         8.5, color=MUTED, after=6)
    p = doc.add_paragraph()
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "2-2" \\h \\z \\u'; r._r.append(it)
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); r._r.append(sep)
    style_run(p.add_run("목차를 표시하려면 문서를 열고 F9로 필드를 업데이트하세요."), 9.5, color=MUTED)
    e = p.add_run()
    ec = OxmlElement("w:fldChar"); ec.set(qn("w:fldCharType"), "end"); e._r.append(ec)
    try:
        st = doc.settings.element
        if st.find(qn("w:updateFields")) is None:
            uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); st.append(uf)
    except Exception:
        pass


# ---- 기업 목록 --------------------------------------------------------------
def build_company_list(doc, ks, demands):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(0); h.paragraph_format.space_after = Pt(2)
    style_run(h.add_run("대상 기업 목록"), 21, bold=True, color=NAVY)
    cap = para(doc, "TARGET COMPANIES", 9, bold=True, color=ACCENT, spacing=50, after=6)
    para_border(cap, "bottom", NAVY, 14, 6)
    para(doc, f"기업 {len(ks)}개사", 10, color=MUTED, after=12)

    t = doc.add_table(rows=1, cols=3)
    table_grid(t, HAIR, 4, "all"); table_cellmar(t)
    for c, txt in zip(t.rows[0].cells, ("번호", "기업명", "기술명")):
        shade_cell(c, HEAD_BG); cell_vcenter(c)
        fill_cell(c, txt, 9.5, bold=True, color=HEAD_FG, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, k in enumerate(ks):
        v = demands[k]
        row = t.add_row().cells
        if idx % 2 == 1:
            for c in row:
                shade_cell(c, ZEBRA)
        fill_cell(row[0], k, 9.5, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(row[1], v["기업명"], 9.5, line=1.25)
        fill_cell(row[2], "\n".join(tech_names(v)), 9, line=1.25)
        for c in row:
            cell_vcenter(c)
    table_fixed(t, [700, 2000, 5940])


# ---- 기업 블록 --------------------------------------------------------------
def build_company(doc, k, dm, pidf):
    h2 = doc.add_heading(level=2)
    h2.paragraph_format.page_break_before = True
    h2.paragraph_format.space_before = Pt(0); h2.paragraph_format.space_after = Pt(5)
    badge_txt = f" No.{k} "
    _off = _text_w(badge_txt, 12) + _text_w("  ", 12)
    h2.paragraph_format.left_indent = Pt(_off); h2.paragraph_format.first_line_indent = Pt(-_off)
    badge = h2.add_run(badge_txt)
    style_run(badge, 12, bold=True, color="FFFFFF"); run_shade(badge, ACCENT)
    style_run(h2.add_run("  "), 12)
    style_run(h2.add_run(dm["기업명"]), 14, bold=True, color=NAVY)
    para_border(h2, "bottom", HAIR, 6, 6)

    # 기업 정보 표 — 기술 정보는 라벨 없이 싣는다(매칭 기준 비노출)
    rows = [("기업명", dm["기업명"], False)]
    names = tech_names(dm)
    if names:
        gubun = " / ".join(x for x in (dm.get("기술유형", ""), dm.get("기술분야", "")) if x)
        rows.append(("기술명", "\n".join(names) + (f"\n({gubun})" if gubun else ""), False))
    for body in tech_bodies(dm):
        rows.append(("기술 내용", body, True))
    if (dm.get("키워드") or "").strip():
        rows.append(("핵심 키워드",
                     " · ".join(x for x in dm["키워드"].split(";") if x), False))

    t = doc.add_table(rows=0, cols=2)
    table_grid(t, HAIR, 4, "all"); table_cellmar(t, 44, 44, 110, 110)
    for label, val, prose in rows:
        cells = t.add_row().cells
        shade_cell(cells[0], LABEL_BG); cell_vcenter(cells[0])
        fill_cell(cells[0], label, 9.5, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[1], val, 9,
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY if prose else None,
                  family=SERIF if prose else gr.SANS, line=1.28)
    table_fixed(t, [1560, 7080])

    # 최종 추천 Top10
    section_label(doc, "최종 추천 과제  Top 10", before=9, after=5)
    # 적합도 점수는 싣지 않는다(순위로만 제시)
    heads = ("순위", "과제명", "수행기관", "수행년도", "특허", "매칭 근거")
    t = doc.add_table(rows=1, cols=len(heads))
    table_grid(t, HAIR, 4, "all"); table_cellmar(t, 32, 32, 70, 70)
    for c, txt in zip(t.rows[0].cells, heads):
        shade_cell(c, HEAD_BG); cell_vcenter(c)
        fill_cell(c, txt, 9, bold=True, color=HEAD_FG, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, tp in enumerate(dm["top10"]):
        cells = t.add_row().cells
        if idx % 2 == 1:
            for c in cells:
                shade_cell(c, ZEBRA)
        fill_cell(cells[0], str(tp["rank"]), 11, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[1], tp["과제명"], 8.6, line=1.16)
        fill_cell(cells[2], tp.get("수행기관", ""), 8.4, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[3], "\n".join(year_lines(tp.get("과제설명문", ""))), 8.4,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        npat = patent_count(tp)
        fill_cell(cells[4], f"{npat}건", 8.4, align=WD_ALIGN_PARAGRAPH.CENTER,
                  bold=npat > 0, color=NAVY if npat > 0 else MUTED)
        fill_cell(cells[5], rename_rnd(tp.get("판단근거", "")), 8.4,
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY, family=SERIF, line=1.16)
        for c in cells:
            cell_vcenter(c)
    table_fixed(t, [520, 2860, 1220, 860, 520, 2660])
    rows_cantsplit(t)

    for tp in dm["top10"]:
        gr.build_top_detail(doc, neutralize(tp), pidf)


def build():
    demands, pidf = load_data()
    try:
        gr.LOGO_PATH = gr.make_top_logo()
    except Exception as e:
        print("로고 크롭 실패(로고 없이 진행):", e)
        gr.LOGO_PATH = None

    doc = Document()
    setup_styles(doc)
    set_margins(doc.sections[0])

    ks = sorted(demands, key=sort_key)
    n_rec = sum(len(demands[k]["top10"]) for k in ks)
    n_proj = len({t["과제고유번호"] for k in ks for t in demands[k]["top10"]})

    build_cover(doc, len(ks), n_rec, n_proj)
    build_intro_toc(doc, demands, n_rec, n_proj)

    doc.add_section(WD_SECTION.NEW_PAGE)
    body = doc.sections[-1]
    set_margins(body)
    set_pgnum_start(body, 1)
    setup_running_header(body)
    body.footer.is_linked_to_previous = False
    fp = body.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gr.add_page_field(fp)
    doc.sections[0].footer.is_linked_to_previous = False
    setup_cover_header(doc.sections[0])

    build_company_list(doc, ks, demands)
    for k in ks:
        build_company(doc, k, demands[k], pidf)

    doc.save(OUT)
    print("saved:", OUT)
    print(f"  기업 {len(ks)}개사 · 추천 {n_rec}건 · 중복제외 과제 {n_proj}개")
    return OUT


if __name__ == "__main__":
    build()
